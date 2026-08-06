#!/usr/bin/env python3
from __future__ import annotations

import csv
import json
import os
import re
import sqlite3
import subprocess
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / 'data' / 'attachments.db'
LOCAL_TZ = ZoneInfo('Asia/Ho_Chi_Minh')

def local_today() -> date:
    return datetime.now(LOCAL_TZ).date()

TZ_TODAY = local_today

FIELD_MAP = {
    'overdraft': {
        'date': ['prvexprdt', 'ngayhethan', 'ngay het han', 'hết hạn', 'het han'],
        'name': ['custnm', 'tenkh', 'ten khach hang', 'tên khách hàng'],
        'seq': ['custseq', 'makh', 'ma khach hang', 'mã khách hàng'],
        'addr': ['addr1', 'diachi', 'dia chi', 'địa chỉ'],
    },
    'guarantee': {
        'date': ['matdt', 'ngayhethan', 'ngay het han', 'hết hạn', 'het han'],
        'trno': ['trno', 'soblanh', 'so bao lanh', 'số bảo lãnh'],
        'cust': ['custno', 'custnm', 'tenkh', 'ten khach hang', 'tên khách hàng'],
    },
}


def norm(s: Any) -> str:
    return re.sub(r'[^a-z0-9]+', '', str(s or '').strip().lower())


def latest_path(statement_type: str) -> Path:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    row = conn.execute(
        'SELECT * FROM statement_uploads WHERE statement_type = ? ORDER BY uploaded_at DESC LIMIT 1',
        (statement_type,),
    ).fetchone()
    conn.close()
    if not row:
        raise SystemExit(f'Chưa có file sao kê {statement_type}')
    p = Path(row['path'])
    if not p.exists():
        raise SystemExit(f'File không tồn tại: {p}')
    return p


def convert_to_csv(path: Path) -> Path:
    if path.suffix.lower() == '.csv':
        return path
    tmp = Path(tempfile.mkdtemp(prefix='statement_csv_'))
    subprocess.run(
        ['libreoffice', '--headless', '--convert-to', 'csv', '--outdir', str(tmp), str(path)],
        check=True,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    out = tmp / (path.stem + '.csv')
    if not out.exists():
        found = list(tmp.glob('*.csv'))
        if not found:
            raise SystemExit('Không convert được Excel sang CSV')
        return found[0]
    return out


def rows_to_dicts(rows: list[list[Any]]) -> list[dict[str, Any]]:
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return []
    header_idx = 0
    for i, r in enumerate(rows[:20]):
        nr = {norm(c) for c in r}
        hits = sum(any(norm(alias) in nr for alias in aliases) for aliases in FIELD_MAP['overdraft'].values())
        hits += sum(any(norm(alias) in nr for alias in aliases) for aliases in FIELD_MAP['guarantee'].values())
        if hits >= 2:
            header_idx = i
            break
    headers = [str(c).strip() for c in rows[header_idx]]
    out = []
    for r in rows[header_idx + 1:]:
        if len(r) < len(headers):
            r = r + [''] * (len(headers) - len(r))
        out.append({headers[i]: (r[i] if i < len(r) else '') for i in range(len(headers))})
    return out


def read_excel_rows(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() == '.xls':
        import xlrd
        book = xlrd.open_workbook(str(path))
        sh = book.sheet_by_index(0)
        rows = []
        for rx in range(sh.nrows):
            vals = []
            for cx, cell in enumerate(sh.row(rx)):
                if cell.ctype == xlrd.XL_CELL_DATE:
                    vals.append(datetime(*xlrd.xldate_as_tuple(float(cell.value), book.datemode)).date().isoformat())
                else:
                    vals.append(cell.value)
            rows.append(vals)
        return rows_to_dicts(rows)
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = [[cell for cell in row] for row in ws.iter_rows(values_only=True)]
        return rows_to_dicts(rows)
    return []


def read_rows(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() in {'.xls', '.xlsx'}:
        try:
            return read_excel_rows(path)
        except Exception as exc:
            # Fallback through LibreOffice for odd workbooks; keep original error if both fail.
            excel_error = exc
        else:
            excel_error = None
    else:
        excel_error = None
    try:
        csv_path = convert_to_csv(path)
    except subprocess.CalledProcessError as exc:
        msg = (exc.stderr or exc.stdout or str(exc)).strip()
        if excel_error:
            raise SystemExit(f'Không đọc được Excel bằng thư viện ({excel_error}) và LibreOffice ({msg})') from exc
        raise SystemExit(f'Không convert được file sao kê bằng LibreOffice: {msg}') from exc
    raw = csv_path.read_bytes()
    for enc in ['utf-8-sig', 'utf-16', 'cp1258', 'cp1252', 'latin1']:
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode('utf-8', errors='replace')
    sample = text[:4096]
    dialect = csv.Sniffer().sniff(sample, delimiters=',;\t') if sample.strip() else csv.excel
    rows = list(csv.reader(text.splitlines(), dialect))
    return rows_to_dicts(rows)


def get_value(row: dict[str, Any], aliases: list[str]) -> str:
    lookup = {norm(k): v for k, v in row.items()}
    for alias in aliases:
        v = lookup.get(norm(alias))
        if v not in (None, ''):
            return str(v).strip()
    return ''


def parse_date(v: Any) -> date | None:
    if v in (None, ''):
        return None
    s = str(v).strip()
    if re.fullmatch(r'\d+(\.0+)?', s):
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(s))).date()
        except Exception:
            pass
    for fmt in ['%d/%m/%Y', '%m/%d/%Y', '%d-%m-%Y', '%m-%d-%Y', '%Y-%m-%d']:
        try:
            return datetime.strptime(s.split()[0], fmt).date()
        except Exception:
            pass
    try:
        return datetime.fromisoformat(s.split()[0]).date()
    except Exception:
        return None


def targets_for(today: date) -> set[date]:
    targets = {today}
    if today.weekday() == 0:
        targets.add(today - timedelta(days=1))
        targets.add(today - timedelta(days=2))
    return targets


def build(statement_type: str, today_s: str | None = None) -> dict[str, Any]:
    today = datetime.strptime(today_s, '%Y-%m-%d').date() if today_s else TZ_TODAY()
    path = latest_path(statement_type)
    rows = read_rows(path)
    fmap = FIELD_MAP[statement_type]
    targets = targets_for(today)
    due = []
    for row in rows:
        dt = parse_date(get_value(row, fmap['date']))
        if not dt or dt not in targets:
            continue
        if statement_type == 'overdraft':
            due.append({
                'custseq': get_value(row, fmap['seq']),
                'custnm': get_value(row, fmap['name']),
                'addr1': get_value(row, fmap['addr']),
                'date': dt.isoformat(),
            })
        else:
            due.append({
                'trno': get_value(row, fmap['trno']),
                'custno': get_value(row, fmap['cust']),
                'date': dt.isoformat(),
            })
    due.sort(key=lambda r: (r['date'], r.get('custnm') or r.get('trno') or ''))
    today_label = today.strftime('%d/%m/%Y')
    if statement_type == 'overdraft':
        lines = [f"{i}. {r['custseq']} - {r['custnm']} - {r['addr1']} - Hết hạn: {datetime.strptime(r['date'], '%Y-%m-%d').strftime('%d/%m/%Y')}" for i, r in enumerate(due, 1)]
        message = (f"Hôm nay ngày {today_label}, có {len(due)} khách hàng thấu chi đến hạn, bao gồm:\n" + '\n'.join(lines)) if due else f"Hôm nay ngày {today_label}, không có khách hàng thấu chi đến hạn."
    else:
        lines = [f"{i}. Số bảo lãnh {r['trno'] or 'N/A'}, tên KH: {r['custno'] or 'N/A'}, ngày hết hạn: {datetime.strptime(r['date'], '%Y-%m-%d').strftime('%d/%m/%Y')}" for i, r in enumerate(due, 1)]
        message = (f"Hôm nay ngày {today_label}, có {len(due)} khách hàng bảo lãnh đến hạn, bao gồm:\n" + '\n'.join(lines)) if due else f"Hôm nay ngày {today_label}, không có khách hàng bảo lãnh đến hạn."
    return {'success': True, 'type': statement_type, 'source': str(path), 'today': today.isoformat(), 'targetDates': sorted(d.isoformat() for d in targets), 'count': len(due), 'dueRows': due, 'message': message}


if __name__ == '__main__':
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument('type', choices=['overdraft', 'guarantee'])
    ap.add_argument('--date')
    args = ap.parse_args()
    print(json.dumps(build(args.type, args.date), ensure_ascii=False))
